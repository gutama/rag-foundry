"""Shared pytest fixtures for RAG Foundry test suite."""

import io
import docx
import fitz
import openpyxl
import pytest


@pytest.fixture
def sample_pdf_bytes():
    """Generate sample digital PDF in bytes with text and headings."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Institutional Architecture Report", fontsize=18)
    page.insert_text(
        (50, 100),
        "This section details the canonical parsing model.",
        fontsize=12,
    )
    pdf_buffer = io.BytesIO()
    doc.save(pdf_buffer)
    doc.close()
    return pdf_buffer.getvalue()


@pytest.fixture
def scanned_pdf_bytes():
    """Generate sample scanned PDF (image-only page with no text layer)."""
    doc = fitz.open()
    page = doc.new_page()
    pix = fitz.Pixmap(fitz.csRGB, fitz.Rect(0, 0, 200, 200), False)
    pix.clear_with(255)
    page.insert_image(page.rect, pixmap=pix)

    pdf_buffer = io.BytesIO()
    doc.save(pdf_buffer)
    doc.close()
    return pdf_buffer.getvalue()


@pytest.fixture
def sample_docx_bytes():
    """Generate sample DOCX in bytes with heading, paragraph, and table."""
    doc = docx.Document()
    doc.add_heading("Financial Summary 2026", level=1)
    doc.add_paragraph("Revenue increased by 15% year over year.")

    table = doc.add_table(rows=2, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Quarter"
    hdr_cells[1].text = "Revenue"

    row_cells = table.rows[1].cells
    row_cells[0].text = "Q1"
    row_cells[1].text = "$1,500,000"

    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)
    return docx_buffer.getvalue()


@pytest.fixture
def sample_xlsx_bytes():
    """Generate sample XLSX in bytes with spreadsheet sheet and table data."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Q1 Data"
    ws.append(["Metric", "Value"])
    ws.append(["Ingestion Throughput", "500 docs/min"])

    xlsx_buffer = io.BytesIO()
    wb.save(xlsx_buffer)
    return xlsx_buffer.getvalue()
